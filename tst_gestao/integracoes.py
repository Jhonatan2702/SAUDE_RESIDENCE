"""Apoio às integrações: chaves de API, webhooks e leitura de planilhas."""
from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import re
import secrets
import threading
import unicodedata
import os
import urllib.request
from datetime import date, datetime
from urllib.error import HTTPError, URLError

from .db import connect

EVENTOS_WEBHOOK = [
    "paciente.criado",
    "paciente.atualizado",
    "terapia.criada",
    "terapia.atualizada",
    "movimentacao.criada",
]
LIMITE_LINHAS = 5000
DADOSBI_URL_PADRAO = "http://35.247.217.138:8081/SaudeResidenceHomeCareServices/ultron/dadosbi/lista/dadospaciente"
DADOSBI_LIMITE_RESPOSTA = 25_000_000
DADOSBI_LIMITE_REGISTROS = 20_000
DADOSBI_INTERVALO = 60 * 60
_dadosbi_trava = threading.Lock()
_dadosbi_execucao_trava = threading.Lock()
_dadosbi_thread_iniciada = False


def url_dadosbi() -> str:
    return os.environ.get("TST_DADOSBI_URL", DADOSBI_URL_PADRAO).strip()


def _data_dadosbi(valor) -> str | None:
    texto = str(valor or "").strip()
    if not texto:
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    raise ValueError("Data inválida na resposta da API.")


def _salvar_status_dadosbi(status: dict) -> None:
    with connect() as conn:
        atual = conn.execute("SELECT id FROM dadosbi_status ORDER BY id").fetchone()
        valores = (
            status.get("ultima_execucao"), status.get("ultima_sucesso"), status.get("status"),
            status.get("pacientes_criados"), status.get("pacientes_atualizados"),
            status.get("visitas_importadas"), status.get("erro"),
        )
        if atual:
            conn.execute(
                "UPDATE dadosbi_status SET ultima_execucao=?, ultima_sucesso=?, status=?, "
                "pacientes_criados=?, pacientes_atualizados=?, visitas_importadas=?, erro=? WHERE id=?",
                (*valores, atual["id"]),
            )
        else:
            conn.execute(
                "INSERT INTO dadosbi_status (ultima_execucao, ultima_sucesso, status, pacientes_criados, "
                "pacientes_atualizados, visitas_importadas, erro) VALUES (?,?,?,?,?,?,?)",
                valores,
            )


def status_dadosbi() -> dict | None:
    with connect() as conn:
        row = conn.execute("SELECT * FROM dadosbi_status ORDER BY id").fetchone()
    return dict(row) if row else None


def _sincronizar_dadosbi() -> dict:
    """Busca pacientes e visitas na API Saúde Residence e importa sem duplicar visitas."""
    agora = datetime.now().isoformat(timespec="seconds")
    try:
        req = urllib.request.Request(url_dadosbi(), headers={"Accept": "application/json", "User-Agent": "TST-GESTAO/1"})
        with urllib.request.urlopen(req, timeout=30) as resposta:
            bruto = resposta.read(DADOSBI_LIMITE_RESPOSTA + 1)
        if len(bruto) > DADOSBI_LIMITE_RESPOSTA:
            raise ValueError("Resposta da API excede o limite permitido.")
        payload = json.loads(bruto.decode("utf-8-sig"))
        registros = payload.get("dados") if isinstance(payload, dict) and payload.get("sucesso") is True else None
        if not isinstance(registros, list):
            raise ValueError("Formato de resposta não reconhecido.")
        if len(registros) > DADOSBI_LIMITE_REGISTROS:
            raise ValueError("Quantidade de registros excede o limite permitido.")

        criados = atualizados = visitas_importadas = erros = 0
        with connect() as conn:
            pacientes_validos = set()
            for registro in registros:
                try:
                    if not isinstance(registro, dict) or registro.get("id") in (None, "") or not str(registro.get("nmpaciente") or "").strip():
                        raise ValueError("Paciente sem identificador ou nome.")
                    codigo = str(registro["id"]).strip()
                    nome = str(registro["nmpaciente"]).strip()
                    admissao = _data_dadosbi(registro.get("dataadmissao"))
                    alta = _data_dadosbi(registro.get("dataalta"))
                    situacao = str(registro.get("status") or "").strip().upper()
                    if situacao == "ALTA" and not alta:
                        raise ValueError("Paciente com alta sem data de alta.")
                    if situacao == "EM ATENDIMENTO":
                        alta = None
                    motivo = str(registro.get("condicaoalta") or "").strip() or None
                    existente = conn.execute("SELECT id FROM pacientes WHERE codigo = ?", (codigo,)).fetchone()
                    valores = (
                        nome,
                        str(registro.get("nmoperadora") or "").strip() or None,
                        str(registro.get("nmservico") or "").strip() or None,
                        admissao,
                        alta,
                        motivo if alta else None,
                    )
                    if existente:
                        conn.execute(
                            "UPDATE pacientes SET nome=?, operadora=COALESCE(?, operadora), "
                            "servico=COALESCE(?, servico), data_admissao=COALESCE(?, data_admissao), "
                            "data_alta=?, motivo_alta=? WHERE codigo=?",
                            (*valores, codigo),
                        )
                        atualizados += 1
                    else:
                        conn.execute(
                            "INSERT INTO pacientes (codigo, nome, operadora, servico, data_admissao, data_alta, motivo_alta) "
                            "VALUES (?,?,?,?,?,?,?)",
                            (codigo, *valores),
                        )
                        criados += 1
                    pacientes_validos.add(codigo)
                except (TypeError, ValueError):
                    erros += 1

            for registro in registros:
                if not isinstance(registro, dict):
                    continue
                codigo = str(registro.get("id") or "").strip()
                if codigo not in pacientes_validos:
                    continue
                visitas = registro.get("visitas") or []
                if not isinstance(visitas, list):
                    erros += 1
                    continue
                for visita in visitas:
                    try:
                        if not isinstance(visita, dict) or visita.get("id") in (None, ""):
                            raise ValueError("Visita sem identificador.")
                        origem_id = str(visita["id"]).strip()
                        data = _data_dadosbi(visita.get("datavisita"))
                        if not data:
                            raise ValueError("Visita sem data.")
                        profissional = str(visita.get("nmprofessional") or "").strip() or None
                        especialidade = str(visita.get("nmespecialidade") or "").strip()
                        texto = "Visita importada da API Saúde Residence."
                        if especialidade:
                            texto += f"\nEspecialidade: {especialidade}"
                        if profissional:
                            texto += f"\nProfissional: {profissional}"
                        existente = conn.execute(
                            "SELECT evolucao_id FROM dadosbi_visitas WHERE origem_id = ?", (origem_id,)
                        ).fetchone()
                        if existente:
                            conn.execute(
                                "UPDATE evolucoes SET codigo_paciente=?, data=?, tipo='Visita', profissional=?, texto=? "
                                "WHERE id=?",
                                (codigo, data, profissional, texto, existente["evolucao_id"]),
                            )
                        else:
                            insercao = conn.execute(
                                "INSERT INTO evolucoes (codigo_paciente, data, tipo, profissional, texto, autor, criado_em) "
                                "VALUES (?,?,?,?,?,?,?)",
                                (codigo, data, "Visita", profissional, texto, "Integração Saúde Residence", agora),
                            )
                            conn.execute(
                                "INSERT INTO dadosbi_visitas (origem_id, codigo_paciente, evolucao_id) VALUES (?,?,?)",
                                (origem_id, codigo, insercao.lastrowid),
                            )
                            visitas_importadas += 1
                    except (TypeError, ValueError):
                        erros += 1

        resultado = {
            "ultima_execucao": agora,
            "ultima_sucesso": agora,
            "status": "sucesso" if not erros else "parcial",
            "pacientes_criados": criados,
            "pacientes_atualizados": atualizados,
            "visitas_importadas": visitas_importadas,
            "erro": f"{erros} registro(s) inválido(s)." if erros else None,
        }
        _salvar_status_dadosbi(resultado)
        return resultado
    except Exception as exc:
        if isinstance(exc, HTTPError):
            erro = f"API retornou HTTP {exc.code}."
        elif isinstance(exc, URLError):
            erro = "Não foi possível conectar à API."
        elif isinstance(exc, (ValueError, json.JSONDecodeError)):
            erro = "Resposta da API inválida ou fora do formato esperado."
        else:
            erro = "Falha ao importar os dados da API."
        anterior = status_dadosbi() or {}
        _salvar_status_dadosbi({
            "ultima_execucao": agora,
            "ultima_sucesso": anterior.get("ultima_sucesso"),
            "status": "erro",
            "pacientes_criados": 0,
            "pacientes_atualizados": 0,
            "visitas_importadas": 0,
            "erro": erro,
        })
        raise RuntimeError(erro) from exc


def sincronizar_dadosbi() -> dict:
    if not _dadosbi_execucao_trava.acquire(blocking=False):
        raise RuntimeError("Já existe uma sincronização Saúde Residence em andamento.")
    try:
        return _sincronizar_dadosbi()
    finally:
        _dadosbi_execucao_trava.release()


def iniciar_sincronizacao_dadosbi(logger) -> None:
    """Inicia uma única rotina automática no processo principal do aplicativo."""
    global _dadosbi_thread_iniciada
    with _dadosbi_trava:
        if _dadosbi_thread_iniciada:
            return
        _dadosbi_thread_iniciada = True

    def executar() -> None:
        while True:
            try:
                sincronizar_dadosbi()
            except Exception:
                logger.exception("Falha na sincronização automática Saúde Residence.")
            threading.Event().wait(DADOSBI_INTERVALO)

    threading.Thread(target=executar, name="sincronizacao-dadosbi", daemon=True).start()


# ---------- chaves ----------
def hash_chave(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def nova_chave() -> tuple[str, str, str]:
    token = "tst_" + secrets.token_urlsafe(32)
    return token, token[:10], hash_chave(token)


# ---------- webhooks ----------
def _enviar(hook, evento: str, dados) -> tuple[bool, str]:
    corpo = json.dumps(
        {"evento": evento, "ocorrido_em": datetime.now().isoformat(timespec="seconds"), "dados": dados},
        ensure_ascii=False,
        default=str,
    ).encode()
    assinatura = hmac.new(hook["segredo"].encode(), corpo, hashlib.sha256).hexdigest()
    req = urllib.request.Request(
        hook["url"],
        data=corpo,
        method="POST",
        headers={
            "Content-Type": "application/json; charset=utf-8",
            "User-Agent": "TST-GESTAO-Webhook/1",
            "X-TST-Event": evento,
            "X-TST-Signature": "sha256=" + assinatura,
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=5) as r:
            ok, resultado = 200 <= r.status < 300, f"HTTP {r.status}"
    except Exception as exc:  # falha de rede/HTTP não pode afetar o sistema
        ok, resultado = False, getattr(exc, "code", None) and f"HTTP {exc.code}" or str(exc)[:200]
    with connect() as conn:
        conn.execute(
            "INSERT INTO webhook_log (webhook_id, evento, ok, resultado, criado_em) VALUES (?,?,?,?,?)",
            (hook["id"], evento, int(ok), resultado, datetime.now().isoformat(timespec="seconds")),
        )
        conn.execute("DELETE FROM webhook_log WHERE id NOT IN (SELECT id FROM webhook_log ORDER BY id DESC LIMIT 300)")
    return ok, resultado


def _enviar_todos(evento: str, dados) -> None:
    with connect() as conn:
        hooks = conn.execute("SELECT * FROM webhooks WHERE ativo = 1").fetchall()
    for h in hooks:
        eventos = (h["eventos"] or "*").split(",")
        if "*" in eventos or evento in eventos:
            _enviar(h, evento, dados)


def notificar(evento: str, dados) -> None:
    threading.Thread(target=_enviar_todos, args=(evento, dados), daemon=True).start()


def testar(hook) -> tuple[bool, str]:
    return _enviar(hook, "teste.ping", {"mensagem": "Teste de webhook do TST-GESTÃO"})


# ---------- planilhas ----------
def norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def _decodificar(dados: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(enc)
        except UnicodeDecodeError:
            continue
    return dados.decode("latin-1")


def ler_planilha(nome: str, dados: bytes) -> list[dict]:
    """Retorna linhas como dict {cabeçalho_normalizado: valor}."""
    nome = (nome or "").lower()
    if nome.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        linhas = list(wb.worksheets[0].iter_rows(values_only=True))
    elif nome.endswith((".csv", ".txt")):
        texto = _decodificar(dados)
        amostra = texto[:4096]
        delim = ";" if amostra.count(";") > amostra.count(",") else ","
        linhas = list(csv.reader(io.StringIO(texto), delimiter=delim))
    else:
        raise ValueError("Formato não suportado. Envie um arquivo .xlsx ou .csv.")
    linhas = [l for l in linhas if l and any(c not in (None, "") for c in l)]
    if len(linhas) < 2:
        raise ValueError("A planilha precisa ter cabeçalho e ao menos uma linha de dados.")
    if len(linhas) - 1 > LIMITE_LINHAS:
        raise ValueError(f"Limite de {LIMITE_LINHAS} linhas por importação excedido.")
    cab = [norm(c) for c in linhas[0]]
    return [{c: v for c, v in zip(cab, l) if c} for l in linhas[1:]]


def valor_data(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", s)
    if m:
        return f"{m[3]}-{int(m[2]):02d}-{int(m[1]):02d}"
    return s[:10] if re.match(r"\d{4}-\d{2}-\d{2}", s) else s


def valor_mes(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m")
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})[/.-](\d{4})", s)
    if m:
        return f"{m[2]}-{int(m[1]):02d}"
    m = re.fullmatch(r"(\d{1,2})[/.-]\d{1,2}[/.-](\d{4})", s)
    if m:
        return f"{m[2]}-{int(m[1]):02d}"
    return s[:7] if re.match(r"\d{4}-\d{2}", s) else s


def valor_texto(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()
